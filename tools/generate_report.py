#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path
from textwrap import dedent

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
FIGURES = ROOT / "docs" / "figures"
OUT.mkdir(exist_ok=True)


def set_font(run, name="Times New Roman", size=14, bold=False, color="000000"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + key))
        if node is None:
            node = OxmlElement("w:" + key)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_table_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def keep_lines_together(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep_lines = OxmlElement("w:keepLines")
    p_pr.append(keep_lines)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_font(run, size=10)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(1.05)
section.right_margin = Inches(0.7)
add_page_number(section.footer.paragraphs[0])

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(14)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.first_line_indent = Inches(0.49)
normal.paragraph_format.space_after = Pt(6)

for style_name, size in (("Title", 20), ("Heading 1", 17), ("Heading 2", 15), ("Heading 3", 14)):
    style = styles[style_name]
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.page_break_before = False


def paragraph(text="", bold_lead=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    if not indent:
        p.paragraph_format.first_line_indent = Inches(0)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def heading(text, level=1):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Inches(0)
    return p


def bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.left_indent = Inches(0.25 + level * 0.25)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_font(r, size=13)
    return p


def numbered(text, number):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{number}. {text}")
    set_font(r, size=13)
    return p


def toc_entry(text, page):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.25), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )
    r = p.add_run(f"{text}\t{page}")
    set_font(r, size=13)
    return p


def table(headers, rows, widths=None, font_size=9.2):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    hdr = tbl.rows[0]
    repeat_table_header(hdr)
    prevent_table_row_split(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = str(text)
        set_cell_shading(cell, "1F4E78")
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                set_font(r, size=font_size, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        row = tbl.add_row()
        prevent_table_row_split(row)
        cells = row.cells
        for i, value in enumerate(values):
            cell = cells[i]
            cell.text = str(value)
            set_cell_shading(cell, "F3F6FA" if row_index % 2 else "FFFFFF")
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Inches(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or str(value).replace(",", ".").replace("−", "-").replace("+", "").replace("%", "").replace(" ", "").replace("—", "").replace(".", "").isdigit() else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    set_font(r, size=font_size)
        if widths:
            for i, width in enumerate(widths):
                cells[i].width = Inches(width)
    if widths:
        for i, width in enumerate(widths):
            hdr.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return tbl


def code_block(code, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.1)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shd)
    lines = dedent(code).strip("\n").splitlines()
    for idx, line in enumerate(lines, 1):
        r = p.add_run(f"{idx:>2}  {line}")
        set_font(r, name="Consolas", size=9.5)
        if idx != len(lines):
            r.add_break()
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Inches(0)
    cap.paragraph_format.space_after = Pt(7)
    r = cap.add_run(caption)
    set_font(r, size=11, bold=True)
    keep_lines_together(p)
    keep_with_next(p)
    return p


def add_figure(path, caption, width=6.6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Inches(0)
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    set_font(r, size=11, bold=True)
    keep_with_next(p)


def page_break():
    doc.add_page_break()


def terminal_image(path):
    lines = [
        "$ dotnet test LibraryDesk.Legacy.sln --collect:\"XPlat Code Coverage\"",
        "LibraryDesk.Legacy -> bin/Debug/net8.0/LibraryDesk.Legacy.dll",
        "LibraryDesk.Legacy.Tests -> bin/Debug/net8.0/LibraryDesk.Legacy.Tests.dll",
        "Test run for LibraryDesk.Legacy.Tests.dll (.NETCoreApp,Version=v8.0)",
        "",
        "Passed!  Failed: 0, Passed: 13, Skipped: 0, Total: 13",
        "Duration: 66 ms",
        "LoanManager.cs line coverage: 74.41%",
    ]
    image = Image.new("RGB", (1700, 520), "#111827")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Menlo.ttc", 34)
    except OSError:
        font = ImageFont.load_default()
    y = 38
    for line in lines:
        color = "#86efac" if line.startswith("Passed") else "#e5e7eb"
        draw.text((42, y), line, font=font, fill=color)
        y += 54
    image.save(path)


# Cover
for text, size, bold in [
    ("МІНІСТЕРСТВО ОСВІТИ І НАУКИ УКРАЇНИ", 13, True),
    ("ВІДОКРЕМЛЕНИЙ СТРУКТУРНИЙ ПІДРОЗДІЛ", 12, True),
    ("ФАХОВИЙ КОЛЕДЖ ІНФОРМАЦІЙНИХ СИСТЕМ І ТЕХНОЛОГІЙ", 12, True),
    ("КИЇВСЬКОГО НАЦІОНАЛЬНОГО ЕКОНОМІЧНОГО УНІВЕРСИТЕТУ", 12, True),
    ("імені ВАДИМА ГЕТЬМАНА", 12, True),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold)

doc.add_paragraph().paragraph_format.space_after = Pt(30)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent = Inches(0)
r = p.add_run("ЗВІТ")
set_font(r, size=21, bold=True)
for text in (
    "із самостійної роботи № 2",
    "з дисципліни Конструювання програмного забезпечення",
    "Аудит якості та рефакторинг успадкованого коду",
):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    r = p.add_run(text)
    set_font(r, size=15, bold="Аудит" in text)

doc.add_paragraph().paragraph_format.space_after = Pt(14)
for text in (
    "Варіант № 2",
    "Проєкт LibraryDesk",
    "Модуль LoanManager Issue",
):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    r = p.add_run(text)
    set_font(r, size=14, bold=True)

doc.add_paragraph().paragraph_format.space_after = Pt(40)
for text in (
    "Виконав студент групи 472",
    "Гордєєв Дмитро Леонідович",
    "Перевірив __________________________",
):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.first_line_indent = Inches(0)
    r = p.add_run(text)
    set_font(r, size=14)

doc.add_paragraph().paragraph_format.space_after = Pt(56)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent = Inches(0)
r = p.add_run("Київ 2026")
set_font(r, size=14)

# Contents and intro
heading("Зміст", 1)
contents = [
    ("Вихідні дані та межі роботи", 2),
    ("1 Інвентаризація та вимірювання", 4),
    ("2 Реєстр дефектів проєктування коду", 7),
    ("3 Страхувальна сітка характеризаційних тестів", 12),
    ("4 План рефакторингу", 13),
    ("5 Виконання рефакторингів", 14),
    ("6 Повторне вимірювання та докази поведінки", 17),
    ("7 Реєстр технічного боргу", 18),
    ("Висновки", 20),
    ("Перелік використаних джерел", 20),
]
for item, page in contents:
    toc_entry(item, page)

heading("Вихідні дані та межі роботи", 1)
paragraph(
    "Робота виконує повний цикл аудиту якості: фіксує початковий стан, "
    "вимірює код, будує страхувальну сітку, планує та виконує малі "
    "рефакторинги, повторно вимірює результат і оцінює залишковий борг. "
    "Головний висновок: максимальна цикломатична складність знижена з 26 "
    "до 10, найдовший метод скорочено з 84 до 15 логічних рядків, а три "
    "копії формули пені замінено одним калькулятором без зміни сценарію."
)
page_break()
table(
    ["Параметр", "Значення"],
    [
        ("Варіант", "2 LibraryDesk"),
        ("Модуль-донор", "LoanManager.Issue"),
        ("Розрахунковий вузол", "строк видачі й пеня за прострочення"),
        ("Джерело", "Б адаптований зріз власного проєкту"),
        ("Обсяг baseline", "551 фізичний LOC, 292 логічний LOC"),
        ("Репозиторій", "https://github.com/kuum-oss/LibraryDesk.Legacy"),
        ("Теги", "baseline, safety-net, after-refactoring"),
    ],
    widths=[1.7, 4.9],
    font_size=10,
)
paragraph(
    "Важливе обмеження: методичні вказівки вимагають код іншої дисципліни "
    "або пакет викладача. Через відсутність пакета використано адаптований "
    "зріз власного LibraryDesk. Це джерело потрібно письмово погодити з "
    "викладачем до подання роботи. У звіті воно не подається як код, "
    "отриманий від викладача."
)
page_break()

# Section 1
heading("1 Інвентаризація та вимірювання", 1)
paragraph(
    "Коміт 9e9e752 з повідомленням baseline вихідний код до аудиту позначено "
    "тегом baseline. На ньому рішення збирається з нульовою кількістю "
    "помилок і попереджень. До вимірювання код не поліпшувався."
)
heading("1.1 Інвентарний опис", 2)
table(
    ["Файл", "Призначення", "Мет.", "Фіз. LOC", "Лог. LOC"],
    [
        ("LegacyModels.cs", "дані читача, книги та видачі", 0, 45, 32),
        ("LoanManager.cs", "видача, строк, пеня, стани", 6, 283, 143),
        ("LoanRepository.cs", "збереження і пошук", 6, 95, 42),
        ("LoanReport.cs", "звіти й картка читача", 4, 86, 44),
        ("Program.cs", "контрольний сценарій", 0, 42, 31),
        ("Разом", "7 типів, 5 файлів", 16, 551, 292),
    ],
    widths=[1.45, 2.9, 0.55, 0.8, 0.8],
)
heading("1.2 Методика і метрики", 2)
paragraph(
    "Скрипт measure_metrics.py однаково обробляє baseline і кінцевий стан. "
    "CC починається з одиниці; кожен if, цикл, case, catch, логічний або "
    "тернарний оператор додає одиницю. Вкладеність рахується від тіла "
    "методу як рівня 1. Fan-out дорівнює кількості інших власних типів, "
    "згаданих класом."
)
metrics_before = [
    ("LoanManager.Issue", 84, 26, 7, 8, 5),
    ("LoanReport.BuildOverdueReport", 19, 10, 4, 3, 3),
    ("LoanManager.CanIssue", 8, 10, 3, 2, 5),
    ("LoanManager.ChangeStatus", 10, 9, 2, 2, 5),
    ("LoanRepository.SumOutstandingFines", 12, 6, 4, 1, 2),
    ("LoanManager.PreviewFine", 9, 5, 3, 2, 5),
    ("LoanReport.BuildReaderCard", 8, 4, 2, 2, 3),
    ("LoanManager.Cancel", 7, 4, 3, 2, 5),
    ("LoanRepository.FindByReader", 5, 4, 3, 1, 2),
    ("LoanRepository.FindByStatus", 5, 3, 2, 1, 2),
    ("LoanRepository.Find", 4, 3, 2, 1, 2),
    ("LoanReport.BuildInventoryReport", 6, 2, 2, 1, 3),
    ("LoanRepository.Clear", 4, 2, 2, 0, 2),
    ("LoanManager.DumpLog", 4, 2, 2, 0, 5),
    ("LoanReport.SaveReport", 4, 2, 2, 2, 3),
]
table(["Метод", "LOC", "CC", "Вкл.", "Парам.", "Fan-out"], metrics_before, widths=[3.25, 0.55, 0.5, 0.55, 0.65, 0.7], font_size=8.6)
page_break()
add_figure(FIGURES / "cc-before.png", "Рисунок 1 Цикломатична складність методів до рефакторингу", width=6.8)
paragraph(
    "Порогом прийнятності обрано CC не більше 8, середню CC не більше 4, "
    "довжину методу до 40 логічних рядків і вкладеність до 3. Issue з CC 26 "
    "потребує щонайменше 26 незалежних шляхів тестування, тому його зміна "
    "без сітки є неприйнятним ризиком. Повна таблиця зберігається у "
    "docs/metrics-before.csv."
)
page_break()

# Section 2
heading("2 Реєстр дефектів проєктування коду", 1)
defects = [
    ("D-01", "LoanManager.cs 14-187", "Довгий метод", 5, 1, 5),
    ("D-02", "три класи", "Дробовик правила пені", 5, 3, 15),
    ("D-03", "Issue 14-22", "Довгий список параметрів", 4, 1, 4),
    ("D-04", "моделі і стани", "Primitive obsession", 4, 4, 16),
    ("D-05", "Issue 26-79", "Вкладеність 7", 5, 1, 5),
    ("D-06", "LoanManager 34-130", "Магічні числа", 4, 4, 16),
    ("D-07", "_lastFine", "Тимчасове поле", 3, 1, 3),
    ("D-08", "два catch", "Проковтнутий виняток", 5, 2, 10),
    ("D-09", "Issue 184-185", "Мертвий код", 2, 1, 2),
    ("D-10", "BuildReaderCard", "Заздрість до даних", 3, 1, 3),
    ("D-11", "LegacyModels", "Неінкапсульовані поля", 4, 4, 16),
    ("D-12", "Issue 20-21", "Flag arguments", 3, 2, 6),
]
table(["ID", "Місце", "Дефект", "С", "Ч", "Пріор."], defects, widths=[0.55, 1.55, 2.75, 0.4, 0.4, 0.65], font_size=8.7)
paragraph(
    "Серйозність оцінюється від 1 до 5. Частота дорівнює кількості місць, "
    "але обмежується чотирма. Пріоритет є добутком цих чисел. Реєстр містить "
    "12 різних типів, а D-02 є профільним дефектом варіанта."
)

cards = [
    ("D-01 Довгий метод", "84 логічні рядки, CC 26, вкладеність 7. Метод поєднує валідацію, строк, пеню, звіт, збереження та сповіщення.", """
public LoanResult Issue(
    Reader? reader,
    List<BookCopy>? books,
    DateTime issuedOn,
    DateTime? returnedOn,
    string subscription,
    bool sendEmail,
    bool printReceipt,
    string operatorName)
""", "R01 Extract Method"),
    ("D-02 Дробовик", "Формула пені повторюється у LoanManager, LoanRepository і LoanReport. Зміна ставки потребує трьох синхронних правок.", """
int overdueDays = (onDate.Date - loan.DueOn.Date).Days;
fine = overdueDays * 2m * loan.Books.Count;
if (loan.Reader != null && loan.Reader.Category == "child")
{
    fine = fine * 0.5m;
}
if (fine > 500m)
{
    fine = 500m;
}
""", "R09 Extract Class і R10 Move Method"),
    ("D-03 Довгий список параметрів", "Вісім позиційних параметрів ускладнюють читання виклику; особливо легко переплутати дві дати й два прапорці.", """
Reader? reader,
List<BookCopy>? books,
DateTime issuedOn,
DateTime? returnedOn,
string subscription,
bool sendEmail,
bool printReceipt,
string operatorName)
""", "R11 Introduce Parameter Object"),
    ("D-04 Одержимість елементарними типами", "Категорія читача, група книги, абонемент і стан подані рядками; будь-яка друкарська помилка компілюється.", """
public string Category = "regular";
public string Group = "regular";
public string Status = "new";
if (subscription == "teacher")
{
    days = 30;
}
loan.Status = "active";
""", "R12 Replace Primitive with enum"),
    ("D-05 Надмірна вкладеність", "Основний сценарій зміщений праворуч сімома рівнями, а причини відмови розпорошені по дереву умов.", """
if (reader != null)
{
    if (!reader.IsBlocked)
    {
        if (books != null)
        {
            if (books.Count > 0)
            {
                if (reader.ActiveLoans < 5)
""", "R06 Guard Clauses"),
    ("D-06 Магічні числа", "Ліміти, строки та ставка пені не мають предметних назв і повторюються.", """
if (reader.ActiveLoans < 5)
int days = 14;
days = 30;
days = 7;
days = 3;
fine = overdueDays * 2m * books.Count;
if (fine > 500m)
""", "R04 Replace Magic Number with Constant"),
    ("D-07 Тимчасове поле", "_lastFine заповнюється лише одним сценарієм і створює прихований стан між викликами.", """
private decimal _lastFine;
// ...
_lastFine = 0m;
if (returnedOn != null)
{
    _lastFine = overdueDays * 2m * books.Count;
    loan.Fine = _lastFine;
}
""", "R15 Replace Temp with Query"),
    ("D-08 Проковтнутий виняток", "Issue повертає успіх навіть після помилки запису. Це справжній дефект поведінки, тому його не виправляють у refactor-коміті.", """
try
{
    _repository.Save(loan);
}
catch
{
    // помилка збереження прихована
}
""", "Окремий fix після завершення аудиту"),
    ("D-09 Мертвий код", "Закоментоване обмеження не виконується і не перевіряється, але створює сумнів щодо чинного правила.", """
result.Success = true;
result.Loan = loan;
result.Receipt = receipt;

// Старий варіант обмежував видачу двома книгами.
// if (books.Count > 2) return new LoanResult();
return result;
""", "R17 Remove Dead Code"),
    ("D-10 Заздрість до чужих даних", "BuildReaderCard читає майже весь Reader і не використовує стан LoanReport.", """
string text = reader.Name.Trim().ToUpperInvariant();
text += " [" + reader.Category + "]";
text += " active=" + reader.ActiveLoans;
text += " fine=" + reader.UnpaidFine.ToString("0.00");
for (int i = 0; i < loans.Count; i++)
{
    if (loans[i].Reader?.Id == reader.Id)
""", "Move Method після поточного циклу"),
    ("D-11 Неінкапсульовані поля", "Публічні поля й List дозволяють обійти будь-який інваріант стану та складу видачі.", """
public class Loan
{
    public int Id;
    public Reader? Reader;
    public List<BookCopy> Books = new();
    public DateTime DueOn;
    public string Status = "new";
    public decimal Fine;
}
""", "R13 Encapsulate Field and Collection"),
    ("D-12 Булеві прапорці", "Два bool утворюють чотири режими, а виклик true false не пояснює намір.", """
bool sendEmail,
bool printReceipt,
// ...
if (sendEmail)
{
    Send(reader);
}
if (printReceipt)
{
    Console.Write(receipt);
}
""", "R19 Remove Flag Argument"),
]

for idx, (title, explanation, code, hypothesis) in enumerate(cards, 1):
    heading(title, 2)
    paragraph(explanation)
    code_block(code, f"Лістинг {idx} Доказ дефекту {title.split()[0]}")
    paragraph(f"Гіпотеза виправлення: {hypothesis}.", bold_lead="Гіпотеза виправлення:")

# Section 3
heading("3 Страхувальна сітка характеризаційних тестів", 1)
paragraph(
    "Сітка створена до першого рефакторингу й зафіксована тегом safety-net. "
    "Вона перевіряє наявну поведінку, а не бажаний дизайн: 8 сценаріїв "
    "строку й пені, 3 помилки та 2 переходи стану. Усі тести мають схему "
    "Метод Умова Результат і структуру arrange act assert."
)
code_block("""
[Fact]
public void Issue_FiveOverdueDaysForTwoBooks_FineIs20()
{
    LoanManager sut = CreateManager();
    LoanResult actual = sut.Issue(
        Request(Reader(), Books(2), new DateTime(2026, 10, 1),
            new DateTime(2026, 10, 20), SubscriptionType.Student));
    Assert.Equal(20m, actual.Loan!.Fine);
    Assert.Equal(LoanStatus.Overdue, actual.Loan.Status);
}
""", "Лістинг 13 Характеризаційний тест пені")
code_block("""
[Fact]
public void Issue_EmptyBookList_ReturnsEmptyError()
{
    LoanManager sut = CreateManager();
    LoanResult actual = sut.Issue(
        Request(Reader(), [], new DateTime(2026, 10, 1), null,
            SubscriptionType.Student));
    Assert.False(actual.Success);
    Assert.Equal("ERR: empty", actual.Error);
}
""", "Лістинг 14 Характеризаційний тест помилки")
code_block("""
[Fact]
public void ChangeStatus_ReturnedToActive_LeavesStateUnchanged()
{
    LoanManager sut = CreateManager();
    Loan loan = new() { Status = LoanStatus.Returned };
    bool actual = sut.ChangeStatus(loan, LoanStatus.Active);
    Assert.False(actual);
    Assert.Equal(LoanStatus.Returned, loan.Status);
}
""", "Лістинг 15 Характеризаційний тест забороненого переходу")

heading("3.1 Шов і мутаційна перевірка", 2)
paragraph(
    "Коміт f9b7560 додає необов'язковий LoanRepository у конструктор "
    "LoanManager і необов'язковий шлях у репозиторій. Значення за "
    "замовчуванням залишають стару поведінку, а тест спрямовує файл до "
    "унікального тимчасового шляху."
)
paragraph(
    "Під час ручної мутації ставка 2m тимчасово замінена на 3m. Два тести "
    "стали червоними: очікуване 20 перетворилося на 30, а дитяча пеня 2 на "
    "3. Після повернення 2m усі 13 тестів знову зелені. Це доводить, що "
    "сітка ловить зміну головного правила."
)
terminal_path = FIGURES / "test-run.png"
terminal_image(terminal_path)
add_figure(terminal_path, "Рисунок 2 Результат запуску характеризаційних тестів", width=6.6)
paragraph("Рядкове покриття LoanManager після завершення становить 74,41 %, гілкове 53,75 %. Ціль 60 % перевиконано.")

# Section 4
heading("4 План рефакторингу", 1)
paragraph("Індекс обчислено як вигода поділена на ризик. За однакового індексу коротший крок стоїть вище.")
plan = [
    (1, "D-06", "R04 константи", 5, 1, 15, "5,00"),
    (2, "D-05", "R06 guard clauses", 5, 1, 20, "5,00"),
    (3, "D-07", "R15 запит", 4, 1, 15, "4,00"),
    (4, "D-01", "R03 змінна", 3, 1, 10, "3,00"),
    (5, "D-02", "R09 FineCalculator", 5, 2, 35, "2,50"),
    (6, "D-01", "R01 методи", 5, 2, 40, "2,50"),
    (7, "D-09", "R17 мертвий код", 2, 1, 5, "2,00"),
    (8, "D-03", "R11 IssueRequest", 4, 2, 35, "2,00"),
    (9, "D-11", "R13 інкапсуляція", 4, 2, 45, "2,00"),
    (10, "D-02", "R10 Move Method", 5, 3, 50, "1,67"),
    (11, "D-04", "R12 enum", 5, 3, 50, "1,67"),
    (12, "D-12", "R19 remove flags", 3, 3, 40, "1,00"),
]
table(["№", "ID", "Рефакторинг", "В", "Р", "хв", "Індекс"], plan, widths=[0.4, 0.55, 2.6, 0.4, 0.4, 0.5, 0.65], font_size=8.8)
heading("4.1 Правила безпеки", 2)
for number, text in enumerate((
    "Один тип рефакторингу дорівнює одному коміту; зміну поведінки не змішувати.",
    "Після кожного кроку запускати build і всі 13 тестів.",
    "Крок довший за 20 хвилин ділити на менші частини.",
    "Якщо тести червоні понад 10 хвилин, відновити файли поточного кроку.",
    "Очікувані значення не змінювати; при зміні сигнатури правити лише виклик.",
    "Перед комітом перевіряти diff, після коміту записувати короткий хеш.",
), 1):
    numbered(text, number)

# Section 5
heading("5 Виконання рефакторингів", 1)
refactors = [
    (1, "R04", "константи", "06e0df5"), (2, "R06", "guard clauses", "8d7d33d"),
    (3, "R15", "запит пені", "abc0150"), (4, "R03", "умова в змінній", "bb1d1e0"),
    (5, "R09", "FineCalculator", "6d44db9"), (6, "R01", "етапи Issue", "e97f88a"),
    (7, "R17", "мертвий код", "c894224"), (8, "R11", "IssueRequest", "fbd9a8b"),
    (9, "R13", "інкапсуляція", "531d7e4"), (10, "R10", "перенесення пені", "1b97e4d"),
    (11, "R12", "enum", "9ac5808"), (12, "R19", "без прапорців", "186691b"),
    (13, "R06", "охоронна гілка звіту", "4e924a6"),
]
table(["№", "Код", "Зміна", "Коміт", "Passed"], [(*x, 13) for x in refactors], widths=[0.45, 0.55, 3.4, 0.9, 0.65], font_size=9)
paragraph(
    "Історія містить 13 refactor-комітів і 12 різних типів. R06 повторено у "
    "другому незалежному місці. Обов'язкові R09 і R10 виконані у комітах "
    "6d44db9 та 1b97e4d. Тег after-refactoring стоїть на останньому "
    "рефакторингу, а документація додана пізніше."
)

heading("5.1 Приклади було і стало", 2)
code_block("""
if (reader != null)
{
    if (!reader.IsBlocked)
    {
        if (books != null)
        {
            if (books.Count > 0)
            {
                // основний сценарій
""", "Лістинг 16 Було вкладене дерево валідації")
code_block("""
if (reader == null) return "ERR: reader";
if (reader.IsBlocked) return "ERR: blocked";
if (books == null) return "ERR: null-books";
if (books.Count == 0) return "ERR: empty";
if (reader.ActiveLoans >= MaximumActiveLoans
    && reader.Category != ReaderCategory.Staff)
    return "ERR: limit";
""", "Лістинг 17 Стало охоронні умови")
code_block("""
fine = overdueDays * 2m * loan.Books.Count;
if (loan.Reader?.Category == "child")
{
    fine *= 0.5m;
}
if (fine > 500m)
{
    fine = 500m;
}
""", "Лістинг 18 Було одна з трьох формул пені")
code_block("""
public decimal Calculate(Loan loan, DateTime onDate)
{
    if (onDate.Date <= loan.DueOn.Date) return 0m;
    int days = (onDate.Date - loan.DueOn.Date).Days;
    decimal fine = days * FinePerBookPerDay * loan.Books.Count;
    if (loan.Reader?.Category == ReaderCategory.Child)
        fine *= ChildFineRate;
    return Math.Min(fine, MaximumFine);
}
""", "Лістинг 19 Стало єдиний FineCalculator")
code_block("""
public LoanResult Issue(
    Reader? reader, List<BookCopy>? books,
    DateTime issuedOn, DateTime? returnedOn,
    string subscription, bool sendEmail,
    bool printReceipt, string operatorName)
""", "Лістинг 20 Було вісім параметрів")
code_block("""
public sealed record IssueRequest(
    Reader? Reader,
    List<BookCopy>? Books,
    DateTime IssuedOn,
    DateTime? ReturnedOn,
    SubscriptionType Subscription,
    string OperatorName);
""", "Лістинг 21 Стало об'єкт параметрів")
code_block("""
LoanResult result = manager.Issue(
    request,
    sendEmail: true,
    printReceipt: false);
""", "Лістинг 22 Було прапорці")
code_block("""
LoanResult result = manager.IssueAndNotify(request);

public LoanResult IssueAndNotify(IssueRequest request)
{
    LoanResult result = Issue(request);
    if (result.Success) SendEmail(request.Reader!);
    return result;
}
""", "Лістинг 23 Стало явний сценарій")
page_break()

# Section 6
heading("6 Повторне вимірювання та докази поведінки", 1)
comparison = [
    ("Логічний LOC", 292, 326, "+34", "+11,64 %"),
    ("Методів", 16, 28, "+12", "+75,00 %"),
    ("Максимальна CC", 26, 10, "−16", "−61,54 %"),
    ("Середня CC", "5,81", "3,04", "−2,77", "−47,68 %"),
    ("Найдовший метод", 84, 15, "−69", "−82,14 %"),
    ("Макс. вкладеність", 7, 3, "−4", "−57,14 %"),
    ("Клони 6+ рядків", 1, 0, "−1", "−100 %"),
    ("Макс. fan-out", 5, 11, "+6", "+120 %"),
    ("Тестів", 0, 13, "+13", "—"),
    ("Покриття LoanManager", "0 %", "74,41 %", "+74,41 п.п.", "—"),
]
table(["Метрика", "До", "Після", "Зміна", "%"], comparison, widths=[2.35, 0.75, 0.75, 0.95, 1.05], font_size=9.2)
paragraph(
    "Максимальна CC лишилася 10 у CanIssue, але знизилася на 61,54 %, що "
    "перевищує допустиму альтернативу 40 %. LOC зріс на 11,64 % через "
    "IssueRequest, FineCalculator і enum. Fan-out зріс до 11, бо нові "
    "типізовані залежності рахуються окремо; це зафіксовано як TD-03."
)
add_figure(FIGURES / "cc-comparison.png", "Рисунок 3 Цикломатична складність до і після", width=6.8)
heading("6.1 Докази незмінності", 2)
numbered("Тести після safety-net змінювалися лише разом із сигнатурами R11, R12 і R19. Числові очікування не змінювалися.", 1)
numbered("На baseline з тестами safety-net Passed 13, Failed 0; на after-refactoring Passed 13, Failed 0.", 2)
numbered("out-before.txt і out-after.txt ідентичні; diff повертає код 0 без рядків відмінності.", 3)
code_block("""
$ git log --oneline safety-net..after-refactoring -- Legacy.Tests
186691b refactor(LoanManager): R19 вилучено прапорці
9ac5808 refactor(models): R12 введено переліки
fbd9a8b refactor(LoanManager): R11 введено IssueRequest

$ git diff --no-index docs/out-before.txt docs/out-after.txt
# виводу немає, код завершення 0
""", "Лістинг 24 Докази історії тестів і однакового сценарію")

# Section 7
heading("7 Реєстр технічного боргу", 1)
paragraph("Тіло оцінено PERT за формулою O плюс 4M плюс P, поділене на 6.")
debt = [
    ("TD-01", "Проковтнуті винятки", "1/2/4", "2,17", "1,00", "2,17"),
    ("TD-02", "Немає тестів сховища і звіту", "2/4/7", "4,17", "1,50", "2,78"),
    ("TD-03", "Fan-out LoanManager 11", "2/3/6", "3,33", "1,00", "3,33"),
    ("TD-04", "CanIssue CC 10", "1,5/2,5/4", "2,58", "0,75", "3,44"),
    ("TD-05", "Пряма залежність від File", "2/4/8", "4,33", "1,25", "3,46"),
    ("TD-06", "Feature envy картки", "1/1,5/3", "1,67", "0,50", "3,34"),
]
table(["ID", "Опис", "O/M/P", "Тіло", "год/міс", "міс"], debt, widths=[0.55, 2.75, 0.9, 0.7, 0.8, 0.6], font_size=8.8)
paragraph("Сумарне тіло боргу 18,25 год. Обслуговування 6,00 год на місяць. Загальна окупність 18,25 ÷ 6,00 = 3,04 місяця.", bold_lead="Сумарне тіло боргу 18,25 год.")
heading("7.1 План погашення", 2)
debt_plan = [
    (1, "TD-01", "не ковтати IOException", "до ЛР-7", "окремий зелений тест"),
    (2, "TD-02", "тести сховища і звітів", "ЛР-7", "покриття ≥ 70 %"),
    (3, "TD-03", "розділити координацію", "до ЛР-8", "fan-out ≤ 7"),
    (4, "TD-06", "ReaderCardPresenter", "тиждень 10", "немає feature envy"),
    (5, "TD-04", "єдине правило CanIssue", "тиждень 10", "CC ≤ 5"),
    (6, "TD-05", "адаптер I/O", "до ЛР-8", "атомарний запис"),
]
table(["№", "ID", "Дія", "Коли", "Критерій"], debt_plan, widths=[0.4, 0.55, 2.55, 1.0, 1.75], font_size=9)
paragraph(
    "TD-01 і TD-02 мають окупність до трьох місяців, тому стоять першими. "
    "TD-01 не виправлено під час рефакторингу навмисно: передавання винятку "
    "змінить зафіксовану поведінку й потребує окремого fix-коміту. Усі шість "
    "пунктів позначені TODO безпосередньо в коді."
)
page_break()

heading("Висновки", 1)
conclusions = [
    "Початковий модуль містив 16 методів, а LoanManager.Issue мав 84 логічні рядки, CC 26 і вкладеність 7.",
    "Сітка з 13 тестів створена до змін, забезпечила понад 60 % покриття й виявила ручну мутацію ставки двома падіннями.",
    "Профільний дробовик усунено обов'язковими Extract Class і Move Method: три формули замінено FineCalculator.Calculate.",
    "Найбільший ефект на витрачений час дали guard clauses, які знизили ризик наступного Extract Method.",
    "Виконано 12 різних типів рефакторингу окремими комітами; на кожному кроці проходили 13 тестів.",
    "Максимальна CC знизилася на 61,54 %, найдовший метод на 82,14 %, вкладеність до 3, а клони до нуля.",
    "Залишковий борг становить 18,25 год і коштує 6 год на місяць; перші дії стосуються винятків і прогалин покриття.",
]
for number, item in enumerate(conclusions, 1):
    numbered(item, number)

heading("Перелік використаних джерел", 1)
sources = [
    "Фаулер М. Рефакторинг Поліпшення наявного коду. Друге видання.",
    "Фезерс М. Ефективна робота з успадкованим кодом.",
    "Microsoft Learn. Документація C# і .NET 8.",
    "xUnit.net. Документація фреймворку модульного тестування.",
    "Coverlet. Документація зі збирання покриття .NET.",
    "Методичні вказівки до самостійної роботи № 2 Аудит якості та рефакторинг успадкованого коду.",
]
for number, source in enumerate(sources, 1):
    numbered(source, number)

doc.core_properties.title = "СР 2 Аудит якості та рефакторинг успадкованого коду"
doc.core_properties.subject = "Варіант 2 LibraryDesk"
doc.core_properties.author = "Гордєєв Дмитро Леонідович"
doc.core_properties.keywords = "LibraryDesk, refactoring, metrics, characterization tests"
doc.core_properties.comments = "Generated from verified repository evidence"

output = OUT / "SR02_Гордєєв_472_2.docx"
doc.save(output)
print(output)
